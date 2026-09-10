"""Final-report data builder + PDF / editable-document generation.

The report is assembled from the *finalised inspection snapshot only* — the
generator never re-runs AI or changes the inspection result. Automated findings
and the inspector's assessment are both preserved (the report reflects the
inspector-confirmed outcome while the audit trail keeps the engine result).
"""
from __future__ import annotations

import datetime as dt
import hashlib
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from .. import models
from ..config import settings
from ..utils import (
    COMPLIANT,
    FINAL_COMPLIANT,
    FINAL_INCONCLUSIVE,
    FINAL_NON_COMPLIANT,
    MANUAL_REVIEW,
    NON_COMPLIANT,
    POTENTIAL_NON_COMPLIANCE,
    human_date,
)
from .rule_engine import ENGINE_VERSION, RULE_BASE_TEXT, aggregate_counts
from .evidence_blocks import (
    auto_evidence_description,
    crop_evidence_image,
    source_image_payload,
    violation_evidence_blocks,
)


def _date(v) -> dt.date:
    if isinstance(v, dt.datetime):
        return v.date()
    return v


def report_payload(db: Session, inspection: models.Inspection) -> dict:
    """Assemble the full report dataset from finalised inspection data."""
    counts = _counts(db, inspection)
    findings = sorted(
        db.query(models.Finding).filter(models.Finding.inspection_id == inspection.id).all(),
        key=lambda f: f.sort_order,
    )
    font_rows = db.query(models.FontAnalysis).filter(
        models.FontAnalysis.inspection_id == inspection.id).all()
    placement_rows = db.query(models.PlacementAnalysis).filter(
        models.PlacementAnalysis.inspection_id == inspection.id).all()
    evidence_rows = db.query(models.Evidence).filter(
        models.Evidence.inspection_id == inspection.id).order_by(models.Evidence.sort_order).all()
    fields = db.query(models.ExtractedField).filter(
        models.ExtractedField.inspection_id == inspection.id,
        models.ExtractedField.status.notin_(["NOT_DETECTED"])).order_by(
        models.ExtractedField.sort).all() if not inspection.id else db.query(models.ExtractedField).filter(
        models.ExtractedField.inspection_id == inspection.id).order_by(models.ExtractedField.sort).all()
    # B3: fall back to rule_results for font/placement summary when analysis rows are empty
    font_rule_results = db.query(models.RuleResult).filter(
        models.RuleResult.inspection_id == inspection.id,
        models.RuleResult.type == "font").all()
    placement_rule_results = db.query(models.RuleResult).filter(
        models.RuleResult.inspection_id == inspection.id,
        models.RuleResult.type == "placement").all()

    confirmed = [f for f in findings if f.inspector_status == "confirmed"
                 or (f.inspector_result == "not_satisfied")]
    rejected = [f for f in findings if f.inspector_status == "rejected"
                or (f.inspector_result == "satisfied")]
    inconclusive = [f for f in findings if f.inspector_status == "inconclusive"
                    or f.inspector_result == "inconclusive"]
    pending = [f for f in findings if f.inspector_status == "pending"]

    snapshot = inspection.rule_set_snapshot or {}

    # Per-violation evidence blocks (label + source image + data-derived
    # description) shared verbatim by the preview, PDF and editable DOCX.
    violation_blocks = violation_evidence_blocks(db, inspection)
    block_by_finding = {b["finding_id"]: b for b in violation_blocks}
    ev_desc_by_id: dict[str, dict] = {}
    for b in violation_blocks:
        for it in b.get("evidence", []):
            ev_desc_by_id[it["evidence_id"]] = {
                "description": it.get("description") or "",
                "source_image": it.get("source_image"),
                "rule_id": it.get("rule_id") or b.get("rule_id"),
                "violation_label": f"{b['finding_id']} — "
                                   f"{b.get('rule_citation') or b.get('title') or ''}",
            }
    imgs_pub = {i.image_id: i for i in inspection.images}

    return {
        "meta": {
            "inspection_id": inspection.inspection_id,
            "report_version": "1.0",
            "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
            "generated_date": dt.datetime.now().strftime("%d %b %Y, %H:%M"),
            "status": "FINALIZED" if inspection.status == "finalized" else inspection.status.upper(),
        },
        "inspection": {
            "inspection_id": inspection.inspection_id,
            "inspection_date": human_date(_date(inspection.inspection_date)),
            "inspection_time": inspection.inspection_date.strftime("%H:%M") if isinstance(inspection.inspection_date, dt.datetime) else "—",
            "inspector": inspection.inspector.username,
            "inspector_name": inspection.inspector.full_name,
            "department": inspection.inspector.department,
            "location": inspection.location,
            "retailer_store": inspection.retailer_store,
            "premises_type": inspection.premises_type,
            "status": inspection.status,
        },
        "product": {
            "product_name": inspection.product_name,
            "brand": inspection.brand,
            "description": inspection.product_description,
            "category": inspection.product_category,
            "package_type": inspection.package_type,
            "fields": [
                {
                    "field_name": f.field_name,
                    "label": f.field_name.replace("_", " ").title(),
                    "value": f.value,
                    "status": f.status,
                    "verified": f.verified,
                    "confidence": f.confidence,
                } for f in fields if f.status != "NOT_DETECTED" and f.value
            ],
        },
        "rule_basis": {
            "base_regulation": snapshot.get("base_regulation", RULE_BASE_TEXT),
            "rule_set_version": snapshot.get("rule_set_version", "—"),
            "effective_date": snapshot.get("effective_date", "—"),
            "amendments": snapshot.get("applicable_amendments", []),
            "engine_version": snapshot.get("engine_version", ENGINE_VERSION),
            "rules_evaluated": counts["evaluated"],
            "coverage_note": snapshot.get("coverage_note", ""),
        },
        "assessment": {
            "automated_result": inspection.automated_result,
            "final_compliance_status": inspection.final_compliance_status,
            "final_remarks": inspection.final_remarks,
        },
        "compliance_summary": counts,
        "findings": {
            "confirmed": [
                {
                    "finding_id": f.finding_id, "requirement": f.requirement,
                    "detected": f.detected_condition, "expected": f.expected_condition,
                    "evidence": f.source_image_id or "", "note": f.inspector_note,
                    "rule": f"{f.rule_number} {f.sub_rule}".strip(),
                    "final": "not_satisfied" if f.inspector_result == "not_satisfied" else "confirmed",
                    "rule_id": (block_by_finding.get(f.finding_id) or {}).get("rule_id"),
                    "category": (block_by_finding.get(f.finding_id) or {}).get("category"),
                    "title": (block_by_finding.get(f.finding_id) or {}).get("title") or f.requirement,
                    "description": (block_by_finding.get(f.finding_id) or {}).get("description")
                                   or f.requirement,
                    "source_image": (block_by_finding.get(f.finding_id) or {}).get("source_image"),
                    "evidence_items": (block_by_finding.get(f.finding_id) or {}).get("evidence", []),
                } for f in confirmed
            ],
            "rejected": [
                {"finding_id": f.finding_id, "requirement": f.requirement,
                 "reason": f.decision_reason, "note": f.inspector_note} for f in rejected
            ],
            "inconclusive": [
                {"finding_id": f.finding_id, "requirement": f.requirement,
                 "reason": f.decision_reason or "Not determinable from available evidence.",
                 "note": f.inspector_note} for f in inconclusive
            ],
            "pending": [{"finding_id": f.finding_id, "requirement": f.requirement} for f in pending],
        },
        "font": {
            # B3 fix: use font_analyses rows when available; fall back to font rule_results
            "analysed": len(font_rows) if font_rows else len(font_rule_results),
            "readable": (
                len([x for x in font_rows if x.readability == "good"])
                if font_rows else
                len([r for r in font_rule_results if r.result == "COMPLIANT"])
            ),
            "review": (
                len([x for x in font_rows if x.manual_review_required])
                if font_rows else
                len([r for r in font_rule_results if r.result == "MANUAL_REVIEW"])
            ),
            "potential": (
                len([x for x in font_rows if x.automated_result == "potential_fail"])
                if font_rows else
                len([r for r in font_rule_results if r.result in ("NON_COMPLIANT", "POTENTIAL_NON_COMPLIANCE")])
            ),
            "note": (
                "Font assessment uses relative height estimates. Physical (mm) measurement "
                "requires a calibration reference and is therefore not claimed automatically."
                if font_rows else
                "Font rule results sourced from rule engine (font_analyses unavailable — "
                "declarations not detected with sufficient spatial data for bbox measurement)."
            ),
            "items": [
                {"field": x.field, "estimate": x.font_estimate, "confidence": x.measurement_confidence,
                 "result": x.automated_result, "inspector_result": x.inspector_result}
                for x in font_rows
            ],
        },
        "placement": {
            # B3 fix: use placement_analyses rows when available; fall back to placement rule_results
            "checked": len(placement_rows) if placement_rows else len(placement_rule_results),
            "passed": (
                len([x for x in placement_rows if x.automated_result == "pass"])
                if placement_rows else
                len([r for r in placement_rule_results if r.result == "COMPLIANT"])
            ),
            "review": (
                len([x for x in placement_rows if x.manual_review_required])
                if placement_rows else
                len([r for r in placement_rule_results if r.result == "MANUAL_REVIEW"])
            ),
            "items": [
                {"field": x.field, "location": x.detected_location, "confidence": x.position_confidence,
                 "result": x.automated_result, "inspector_result": x.inspector_result}
                for x in placement_rows
            ],
        },
        "evidence": {
            "items": len(evidence_rows),
            "rows": [
                {
                    "evidence_id": e.evidence_id, "source_image": e.source_image_id or "",
                    "finding": e.finding_id or "", "description": e.description, "status": e.status,
                    "auto_description": (ev_desc_by_id.get(e.evidence_id) or {}).get("description")
                                        or e.description,
                    "rule_id": (ev_desc_by_id.get(e.evidence_id) or {}).get("rule_id") or e.rule_id,
                    "violation_label": (ev_desc_by_id.get(e.evidence_id) or {}).get("violation_label") or "",
                    "source_image_full": (ev_desc_by_id.get(e.evidence_id) or {}).get("source_image")
                                         or source_image_payload(imgs_pub.get(e.source_image_id or "")),
                } for e in evidence_rows
            ],
            "violations": violation_blocks,
        },
        "observations": {
            "inspector_observation": inspection.inspector_observation,
            "final_remarks": inspection.final_remarks,
            "reviewed_by": inspection.reviewed_by or "",
            "reviewed_at": human_date(inspection.reviewed_at) if inspection.reviewed_at else "—",
        },
    }


def _counts(db: Session, inspection: models.Inspection) -> dict:
    """Inspector-aware compliance summary for reports.

    Uses aggregate_counts_with_overrides so that inspector decisions
    (satisfied / not_satisfied / inconclusive) are reflected in the counts
    shown in Section 6 of the report. The raw rule_results.result is never
    modified — the override is purely a view-level computation.
    """
    from .rule_engine import aggregate_counts_with_overrides
    persisted = db.query(models.RuleResult).filter(
        models.RuleResult.inspection_id == inspection.id).all()
    return aggregate_counts_with_overrides(db, inspection.id, persisted)



def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def _findings_table(findings: list[dict], styles) -> list:
    """Render confirmed / review finding rows as a ReportLab table."""
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Table, TableStyle

    if not findings:
        return [Paragraph("None recorded.", styles["Small"])]

    header = ["Finding", "Requirement / Rule", "Detected", "Expected / Evidence", "Status"]
    rows = [[Paragraph(h, styles["Small"]) for h in header]]
    for f in findings:
        rule = f.get("rule") or ""
        req = f.get("requirement") or ""
        head = (f"<b>{rule}</b> — {req}" if rule else req) or "—"
        note = f.get("note") or ""
        expected = f.get("expected") or "—"
        if note:
            expected = f"{expected} <i>(Inspector note: {note})</i>"
        rows.append([
            Paragraph(f.get("finding_id", "—"), styles["Small"]),
            Paragraph(head, styles["Small"]),
            Paragraph(f.get("detected") or "—", styles["Small"]),
            Paragraph(expected, styles["Small"]),
            Paragraph(str(f.get("final", "CONFIRMED")).replace("_", " ").upper(), styles["Small"]),
        ])
    return [Table(rows, colWidths=[18 * mm, 48 * mm, 34 * mm, 46 * mm, 24 * mm],
                  repeatRows=1,
                  style=TableStyle([
                      ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                      ("FONTSIZE", (0, 0), (-1, -1), 8),
                      ("VALIGN", (0, 0), (-1, -1), "TOP"),
                      ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef4")),
                      ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7fafc")]),
                  ]))]


def _violation_evidence_flowables(data: dict, styles) -> list:
    """Per-violation evidence blocks for the PDF: Evidence-for label + source
    image name + data-derived description + embedded cropped image.

    One block per confirmed violation — never a shared generic gallery.
    """
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, KeepTogether, Paragraph, Spacer, Table, TableStyle

    confirmed_ids = {f["finding_id"] for f in data["findings"]["confirmed"]}
    flow: list = []
    for b in data["evidence"]["violations"]:
        if b["finding_id"] not in confirmed_ids:
            continue
        label = b.get("rule_id") or b.get("rule_citation") or b.get("finding_id")
        title = b.get("title") or b.get("requirement") or ""
        source = b.get("source_image") or (b.get("evidence") or [{}])[0].get("source_image")
        lines: list[Any] = [Paragraph(f"Evidence for: <b>{label}</b> — {title}", styles["Small"])]
        if source:
            lines.append(Paragraph(
                f"Source image: <b>{source.get('image_id', '—')}</b> ({source.get('side', '—')})",
                styles["Small"]))
        if b.get("description"):
            lines.append(Paragraph(f"Description: {b['description']}", styles["Small"]))

        for it in (b.get("evidence") or [])[:2]:
            si = it.get("source_image")
            if not si or not si.get("storage_path"):
                continue
            buf = crop_evidence_image(si["storage_path"], si.get("bbox_px") or None,
                                      it.get("bbox_normalized"))
            if buf is None:
                continue
            cw, ch = (si.get("width") or 1), (si.get("height") or 1)
            px = si.get("bbox_px")
            if px and len(px) == 4:
                cw, ch = max(1, px[2] - px[0]), max(1, px[3] - px[1])
            elif it.get("bbox_normalized"):
                bn = it["bbox_normalized"]
                cw = max(1, int(bn.get("width", 1) * (si.get("width") or 1)))
                ch = max(1, int(bn.get("height", 1) * (si.get("height") or 1)))
            img_h = 70 * mm * (ch / cw)
            if not 0 < img_h <= 125 * mm:
                img_h = 70 * mm
            img_cell = Table([[Image(buf, width=70 * mm, height=img_h)]],
                             colWidths=[78 * mm],
                             style=TableStyle([
                                 ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#9aa7b5")),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                 ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                                 ("TOPPADDING", (0, 0), (-1, -1), 6),
                                 ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                             ]))
            caption = f"Evidence item: {it.get('evidence_id', '—')}"
            if it.get("field_name"):
                caption += f" · {it.get('field_name')}"
            lines.append(Table([[img_cell], [Paragraph(caption, styles["Small"])]],
                               colWidths=[86 * mm]))
        if len(lines) > 1:
            flow.append(KeepTogether(lines))
            flow.append(Spacer(1, 8))
    return flow


def generate_pdf(db: Session, inspection: models.Inspection, out_path: Path) -> str:
    """Render the PDF report (ReportLab platypus). Returns file hash."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    data = report_payload(db, inspection)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="LMHeader", parent=styles["Title"], fontSize=16, alignment=1,
                              textColor=colors.HexColor("#1a3a5c"), spaceAfter=2))
    styles.add(ParagraphStyle(name="LMSub", parent=styles["Normal"], fontSize=9, alignment=1,
                              textColor=colors.grey, spaceAfter=14))
    styles.add(ParagraphStyle(name="H2", parent=styles["Heading2"], fontSize=12, spaceBefore=12,
                              spaceAfter=4, textColor=colors.HexColor("#1a3a5c")))
    styles.add(ParagraphStyle(name="Small", parent=styles["Normal"], fontSize=8.5, leading=11))

    story = []
    story.append(Paragraph("⚖ LEGAL METROLOGY COMPLIANCE SYSTEM", styles["LMHeader"]))
    story.append(Paragraph("INSPECTION REPORT", styles["LMHeader"]))
    story.append(Paragraph(
        f"Inspection ID: {data['inspection']['inspection_id']} &nbsp;•&nbsp; "
        f"Inspection Date: {data['inspection']['inspection_date']} &nbsp;•&nbsp; "
        f"Generated: {data['meta']['generated_date']}", styles["LMSub"]))

    def section(title: str, body: list):
        story.append(Paragraph(title, styles["H2"]))
        story.extend(body)
        story.append(Spacer(1, 4))

    def kv(rows: list[tuple[str, str]]) -> list:
        return [Table([[Paragraph(k, styles["Small"]), Paragraph(str(v) or "—", styles["Small"])]
                       for k, v in rows],
                      colWidths=[55 * mm, 115 * mm],
                      style=TableStyle([
                          ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                          ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                          ("VALIGN", (0, 0), (-1, -1), "TOP"),
                          ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                          ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                      ]))]

    # 1 Inspection information
    insp = data["inspection"]
    section("1. INSPECTION INFORMATION", kv([
        ("Inspection ID", insp["inspection_id"]),
        ("Inspection Date / Time", f"{insp['inspection_date']} {insp['inspection_time']}"),
        ("Inspector", f"{insp['inspector_name'] or insp['inspector']} ({insp['inspector']})"),
        ("Department", insp["department"]),
        ("Location", insp["location"]),
        ("Retailer / Store", insp["retailer_store"]),
        ("Inspection Status", insp["status"].upper()),
    ]))

    # 2 Product information
    prod = data["product"]
    rows = [("Product Name", prod["product_name"]), ("Brand", prod["brand"]),
            ("Description", prod["description"]), ("Category", prod["category"])]
    for f in prod["fields"]:
        if f["value"]:
            rows.append((f["label"], f"{f['value']}  (conf {f['confidence']:.0%})" if f["confidence"] else f["value"]))
    section("2. PRODUCT / PACKAGE INFORMATION", kv(rows))

    # 3 Evidence images (first two)
    images = [i for i in inspection.images if i.side in ("front", "back")][:2]
    if images:
        story.append(Paragraph("3. PACKAGE EVIDENCE", styles["H2"]))
        img_cells = []
        for i in images:
            path = settings.storage_dir / i.storage_path
            if path.exists():
                img_cells.append(Image(str(path), width=70 * mm, height=85 * mm))
                img_cells.append(Paragraph(f"{i.side.title()} — {i.image_id}", styles["Small"]))
        if img_cells:
            story.append(Table([img_cells[:2]], colWidths=[75 * mm, 75 * mm]))
            story.append(Spacer(1, 6))

    # 4 Applicable rule basis
    rb = data["rule_basis"]
    amend = ", ".join(a["name"] for a in rb["amendments"]) or "None applicable"
    section("4. APPLICABLE LEGAL / RULE BASIS", kv([
        ("Base Regulation", rb["base_regulation"]),
        ("Rule Set Version", rb["rule_set_version"]),
        ("Effective Date", rb["effective_date"]),
        ("Applicable Amendments", amend),
        ("Rule Engine Configuration", rb["engine_version"]),
        ("Rules Evaluated", rb["rules_evaluated"]),
    ]))

    # 5 Final assessment
    asm = data["assessment"]
    section("5. FINAL INSPECTION ASSESSMENT", [
        Table([[Paragraph(f"<font size=15 color=white><b>{str(asm['final_compliance_status'] or asm['automated_result'] or 'PENDING').upper()}</b></font>", styles["Small"])]],
              colWidths=[170 * mm],
              style=TableStyle([
                  ("BACKGROUND", (0, 0), (-1, -1),
                   colors.HexColor("#1c7c3c" if asm["final_compliance_status"] == FINAL_COMPLIANT else
                                   ("#b3261e" if asm["final_compliance_status"] == FINAL_NON_COMPLIANT else
                                    ("#8a6d00" if asm["final_compliance_status"] == FINAL_INCONCLUSIVE else
                                     ("#6d28d9" if asm["final_compliance_status"] == "POTENTIAL_NON_COMPLIANCE" else "#5b5b5b"))))),
                  ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                  ("TOPPADDING", (0, 0), (-1, -1), 8),
                  ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
              ])),
        Spacer(1, 4),
        Paragraph("Automated (rule-engine) result: "
                  f"<b>{str(asm['automated_result'] or '—').upper()}</b> — "
                  "final assessment recorded by the inspector.", styles["Small"]),
    ])

    # 6 Compliance summary
    cs = data["compliance_summary"]
    section("6. COMPLIANCE SUMMARY", kv([
        ("Rules Checked", cs["rules_checked"]), ("Compliant", cs["compliant"]),
        ("Non-Compliant", cs["non_compliant"]), ("Manual Review", cs["manual_review"]),
        ("Potential Non-Compliance", cs["potential_non_compliance"]),
    ]))

    # 7 Confirmed findings
    section("7. CONFIRMED FINDINGS", _findings_table(data["findings"]["confirmed"], styles))

    # 7A Per-violation evidence blocks (labelled + source image + description)
    ev_flow = _violation_evidence_flowables(data, styles)
    if ev_flow:
        section("7A. VIOLATION EVIDENCE", ev_flow)
    if data["findings"]["rejected"] or data["findings"]["inconclusive"] or data["findings"]["pending"]:
        section("8. REVIEW / INCONCLUSIVE ITEMS", _findings_table(
            [dict(x, final="INCONCLUSIVE" if x in data["findings"]["inconclusive"] else "NOT CONFIRMED")
             for x in data["findings"]["rejected"] + data["findings"]["inconclusive"] + data["findings"]["pending"]],
            styles))

    # 9 Font & readability summary
    ft = data["font"]
    section("9. FONT & READABILITY ASSESSMENT", kv([
        ("Declarations Analysed", ft["analysed"]),
        ("Readability Acceptable", ft["readable"]),
        ("Manual Review Required", ft["review"]),
        ("Potential Issues", ft["potential"]),
    ]) + [Paragraph(ft["note"], styles["Small"])])

    # 10 Placement summary
    pl = data["placement"]
    section("10. PLACEMENT / FORMAT ASSESSMENT", kv([
        ("Declarations Checked", pl["checked"]), ("Passed", pl["passed"]),
        ("Review Required", pl["review"]),
    ]))

    # 11 Evidence summary + reference table
    ev = data["evidence"]
    ev_tbl = [[Paragraph("Evidence", styles["Small"]), Paragraph("Source", styles["Small"]),
               Paragraph("Finding", styles["Small"]), Paragraph("Status", styles["Small"])]]
    for r in ev["rows"][:16]:
        ev_tbl.append([Paragraph(r["evidence_id"], styles["Small"]), Paragraph(r["source_image"], styles["Small"]),
                       Paragraph(r["finding"], styles["Small"]), Paragraph(r["status"], styles["Small"])])
    section("11. EVIDENCE SUMMARY", [
        Table(ev_tbl, colWidths=[28 * mm, 30 * mm, 26 * mm, 30 * mm],
              style=TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                                ("FONTSIZE", (0, 0), (-1, -1), 8)]))])

    # 12 Inspector observations
    obs = data["observations"]
    section("12. INSPECTOR OBSERVATIONS", [
        Paragraph((obs["inspector_observation"] or "—").replace("\n", "<br/>"), styles["Small"]),
        Spacer(1, 3),
        Paragraph(f"Final remarks: {obs['final_remarks'] or '—'}", styles["Small"]),
    ])

    story.append(Paragraph("13. INSPECTOR REVIEW", styles["H2"]))
    story.extend(kv([("Reviewed By", obs["reviewed_by"]), ("Reviewed Date", obs["reviewed_at"])]))

    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "LEGAL METROLOGY COMPLIANCE SYSTEM — AI-assisted inspection • Rule-based validation • "
        "Evidence-based reporting", styles["Small"]))

    doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                            leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm)
    doc.build(story)
    return _sha256(out_path)


def generate_docx(db: Session, inspection: models.Inspection, out_path: Path) -> str:
    """Generate an editable (DOCX) report from the same finalised snapshot."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Inches, Pt, RGBColor

    data = report_payload(db, inspection)
    doc = Document()
    styles = doc.styles
    title: Any = styles["Title"]
    title.font.size = Pt(18)
    title.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)

    h2: Any = styles["Heading 2"]
    h2.font.size = Pt(12)
    h2.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)

    doc.add_heading("LEGAL METROLOGY COMPLIANCE SYSTEM", level=0)
    doc.add_heading("INSPECTION REPORT", level=1)
    p = doc.add_paragraph(
        f"Inspection ID: {data['inspection']['inspection_id']}    "
        f"Inspection Date: {data['inspection']['inspection_date']}    "
        f"Generated: {data['meta']['generated_date']}")
    p.runs[0].font.size = Pt(9)

    def kv_section(num_title: str, rows: list[tuple[str, str]]):
        doc.add_heading(num_title, level=2)
        table = doc.add_table(rows=len(rows), cols=2)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (k, v) in enumerate(rows):
            table.cell(i, 0).text = k
            table.cell(i, 1).text = str(v) if v else "—"
            for cell in (table.cell(i, 0), table.cell(i, 1)):
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9)
        doc.add_paragraph()

    insp = data["inspection"]
    kv_section("1. INSPECTION INFORMATION", [
        ("Inspection ID", insp["inspection_id"]),
        ("Inspection Date", f"{insp['inspection_date']} {insp['inspection_time']}"),
        ("Inspector", insp["inspector_name"] or insp["inspector"]),
        ("Department", insp["department"]), ("Location", insp["location"]),
        ("Retailer / Store", insp["retailer_store"]), ("Status", insp["status"].upper()),
    ])

    prod = data["product"]
    rows = [("Product Name", prod["product_name"]), ("Brand", prod["brand"]),
            ("Category", prod["category"]), ("Package Type", prod["package_type"])]
    for f in prod["fields"]:
        if f["value"]:
            rows.append((f["label"], f["value"]))
    kv_section("2. PRODUCT / PACKAGE INFORMATION", rows)

    rb = data["rule_basis"]
    amend = ", ".join(a["name"] for a in rb["amendments"]) or "None applicable"
    kv_section("4. APPLICABLE LEGAL / RULE BASIS", [
        ("Base Regulation", rb["base_regulation"]), ("Rule Set Version", rb["rule_set_version"]),
        ("Effective Date", rb["effective_date"]), ("Amendments", amend),
        ("Engine", rb["engine_version"]), ("Rules Evaluated", rb["rules_evaluated"]),
    ])

    asm = data["assessment"]
    kv_section("5. FINAL INSPECTION ASSESSMENT", [
        ("Final Assessment", str(asm["final_compliance_status"] or "PENDING").upper()),
        ("Automated Result", str(asm["automated_result"] or "—").upper()),
        ("Final Remarks", asm["final_remarks"] or "—"),
    ])

    cs = data["compliance_summary"]
    kv_section("6. COMPLIANCE SUMMARY", [
        ("Rules Checked", cs["rules_checked"]), ("Compliant", cs["compliant"]),
        ("Non-Compliant", cs["non_compliant"]), ("Manual Review", cs["manual_review"]),
        ("Potential Non-Compliance", cs["potential_non_compliance"]),
    ])

    confirmed = data["findings"]["confirmed"]
    doc.add_heading("7. CONFIRMED FINDINGS", level=2)
    if not confirmed:
        doc.add_paragraph("No confirmed findings.")
    for cf in confirmed:
        doc.add_paragraph(f"{cf['finding_id']} — {cf['requirement']}", style="List Bullet")
        doc.add_paragraph(f"Observed: {cf['detected']} | Evidence: {cf['evidence']}")

    # 7A Per-violation evidence blocks (label + source image + description)
    from io import BytesIO as _BytesIO

    confirmed_ids = {f["finding_id"] for f in confirmed}
    rendered = 0
    for b in data["evidence"]["violations"]:
        if b["finding_id"] not in confirmed_ids:
            continue
        if rendered == 0:
            doc.add_heading("7A. VIOLATION EVIDENCE", level=2)
        rendered += 1
        label = b.get("rule_id") or b.get("rule_citation") or b.get("finding_id")
        doc.add_paragraph(f"Evidence for: {label} — {b.get('title') or b.get('requirement') or ''}")
        si = b.get("source_image") or (b.get("evidence") or [{}])[0].get("source_image")
        if si:
            doc.add_paragraph(f"Source image: {si.get('image_id', '—')} ({si.get('side', '—')})")
        if b.get("description"):
            doc.add_paragraph(f"Description: {b['description']}")
        for it in (b.get("evidence") or [])[:2]:
            ii = it.get("source_image")
            if not ii or not ii.get("storage_path"):
                continue
            buf = crop_evidence_image(ii["storage_path"], ii.get("bbox_px") or None,
                                      it.get("bbox_normalized"))
            if buf is None:
                continue
            doc.add_picture(_BytesIO(buf.read()), width=Inches(4.2))
            caption = f"Evidence item: {it.get('evidence_id', '—')}"
            if it.get("field_name"):
                caption += f" · {it.get('field_name')}"
            doc.add_paragraph(caption)

    ft = data["font"]
    kv_section("9. FONT & READABILITY ASSESSMENT", [
        ("Declarations Analysed", ft["analysed"]), ("Readable", ft["readable"]),
        ("Review Required", ft["review"]), ("Note", ft["note"]),
    ])

    obs = data["observations"]
    doc.add_heading("12. INSPECTOR OBSERVATIONS", level=2)
    doc.add_paragraph(obs["inspector_observation"] or "—")
    doc.add_paragraph(f"Final remarks: {obs['final_remarks'] or '—'}")
    doc.add_paragraph(f"Reviewed By: {obs['reviewed_by']}    Reviewed At: {obs['reviewed_at']}")

    doc.save(str(out_path))
    return _sha256(out_path)
